"""Link detected invoices into an externally-supplied GSTR-2A return.

For each B2B invoice row in the return, we find the matching PDF from a scan run,
copy it into a single flat folder under a unique, human-meaningful name
(`<GSTIN>__<invoice-no>.pdf`), and write that name into a new `Invoice Ref` column
on the sheet — so a reviewer can jump from any row straight to the source PDF.

Matching keys on `(GSTIN, normalized invoice number)` taken from the PDF's content
and filename — never the folder's FY/month — so an invoice paid in a later year still
links to its row. Rows we cannot match are called out explicitly (`NOT FOUND`,
highlighted, and listed on a `Link Report` sheet). The input template is never
modified; a new `<stem>_linked.xlsx` is written into the run dir.

**The module is split in three on purpose:**

``read_b2b_rows``  parse the return once (the only slow part — an xlsx load).
``match``          pure, in-memory, no I/O. Cheap enough to re-run after *every*
                   review edit, which is what lets the review UI say "you just
                   fixed this GSTIN, and row 143 now matches".
``write_linked``   the expensive half — copy PDFs, write the workbook. Deferred
                   to "Finish & export" so it reflects human corrections.

Only ``write_linked`` touches the disk, and it pulls bytes through a
:class:`FileSource`, so a Drive-hosted invoice is downloaded rather than assumed
to exist at ``doc.path`` (on a Drive run ``doc.path`` is a display string, not a file).
"""
from __future__ import annotations

import re
import shutil
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, Optional

from openpyxl import load_workbook

from ..core.interfaces import FileSource
from ..core.models import Document, LinkStatus, RunResult
from ..io.excel import _FLAG_FILL, _HEADER_FILL, _HEADER_FONT, _autosize, _style_header
from ..io.runstore import RunStore
from ..io.sources import LocalFileSource
from ..matching import vendor

DEFAULT_SHEET = "Part A - B2B Invoices"
DEFAULT_REF_HEADER = "Invoice Ref"
GSTIN_HEADER = "GSTIN of Supplier"
INVNO_HEADER = "Invoice Number"
REPORT_SHEET = "Link Report"

# Row outcomes.
MATCHED = "matched"
NOT_FOUND = "not_found"
AMBIGUOUS = "ambiguous"


# --------------------------------------------------------------------------- #
# data
# --------------------------------------------------------------------------- #
@dataclass
class B2BRow:
    """One B2B invoice row of the supplied return, as read off the sheet."""
    row: int                  # 1-based row number in the sheet
    gstin: str                # raw, as printed
    invoice_no: str           # raw, as printed
    supplier: str = ""        # display only
    value: str = ""           # display only


@dataclass
class RowMatch:
    """What we resolved a :class:`B2BRow` to."""
    row: int
    gstin: str
    invoice_no: str
    status: str                            # matched | not_found | ambiguous
    doc_id: Optional[str] = None           # the invoice satisfying this row
    candidates: list[str] = field(default_factory=list)   # doc ids, when ambiguous
    manual: bool = False                   # bound by a human, not by the matcher
    supplier: str = ""
    value: str = ""

    @property
    def unresolved(self) -> bool:
        return self.status != MATCHED


@dataclass
class LinkPlan:
    """The full two-sided picture: every row, and every invoice's fate."""
    rows: list[RowMatch] = field(default_factory=list)
    by_doc: dict[str, RowMatch] = field(default_factory=dict)   # doc.id -> its row
    unreferenced: list[str] = field(default_factory=list)       # doc ids with no row
    duplicate_filings: int = 0   # rows resolved from >1 copy of the same invoice

    def counts(self) -> dict:
        return {
            "total": len(self.rows),
            "matched": sum(1 for r in self.rows if r.status == MATCHED),
            "not_found": sum(1 for r in self.rows if r.status == NOT_FOUND),
            "ambiguous": sum(1 for r in self.rows if r.status == AMBIGUOUS),
            "unreferenced": len(self.unreferenced),
            "duplicate_filings": self.duplicate_filings,
        }

    def unresolved_rows(self) -> list[RowMatch]:
        """not_found + ambiguous — the queue a human actually has to work."""
        return [r for r in self.rows if r.unresolved]

    def to_dict(self) -> dict:
        return {
            "rows": [r.__dict__ for r in self.rows],
            "unreferenced": list(self.unreferenced),
            "duplicate_filings": self.duplicate_filings,
        }

    @classmethod
    def from_dict(cls, data: dict) -> "LinkPlan":
        rows = [RowMatch(**r) for r in data.get("rows", [])]
        plan = cls(rows=rows,
                   unreferenced=list(data.get("unreferenced", [])),
                   duplicate_filings=int(data.get("duplicate_filings", 0)))
        plan.by_doc = {r.doc_id: r for r in rows if r.doc_id}
        return plan


@dataclass
class LinkReport:
    """Result of actually writing the linked workbook + flat folder."""
    total: int = 0
    matched: int = 0
    not_found: int = 0
    ambiguous: int = 0
    copy_failed: int = 0
    duplicate_filings: int = 0
    unreferenced_invoices: int = 0
    unmatched_rows: list = field(default_factory=list)   # (row, gstin, invno)
    ambiguous_rows: list = field(default_factory=list)   # (row, gstin, invno, n)
    out_path: Path | None = None
    flat_dir: Path | None = None


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _sanitize(s: str) -> str:
    """Filesystem-safe token: slashes/spaces -> '-', drop other unsafe chars."""
    s = s.strip().replace("/", "-").replace("\\", "-")
    s = re.sub(r"\s+", "-", s)
    s = re.sub(r"[^A-Za-z0-9._-]", "", s)
    return s.strip("-._")


def _flat_name(gstin_raw: str, invno_raw: str) -> str:
    g = _sanitize(gstin_raw) or "NOGSTIN"
    n = _sanitize(invno_raw) or "NOINV"
    return f"{g}__{n}.pdf"


def _build_index(invoices: Iterable[Document]):
    """Index detected invoices by (norm_gstin, norm_invno) and by norm_invno alone.

    Each doc registers under BOTH its filename id and its printed content number, so a
    junk-named file still matches on the number printed inside the PDF. Inner dicts are
    keyed by doc.id to dedupe.
    """
    by_key: dict[tuple[str, str], dict[str, Document]] = {}
    by_invno: dict[str, dict[str, Document]] = {}
    for d in invoices:
        g = vendor.norm_gstin(d.fields.vendor_gstin)
        for raw in (d.fields.invoice_id, d.fields.invoice_no_content):
            n = vendor.norm_id(raw)
            if not n:
                continue
            by_invno.setdefault(n, {})[d.id] = d
            if g:
                by_key.setdefault((g, n), {})[d.id] = d
    return by_key, by_invno


def _resolve(by_key, by_invno, g: str, n: str) -> list[Document]:
    """Candidate docs for a row: composite key first, then invoice-no-only
    (narrowed by GSTIN when that leaves any candidate)."""
    if not n:
        return []
    if g and (g, n) in by_key:
        return list(by_key[(g, n)].values())
    if n in by_invno:
        cands = list(by_invno[n].values())
        if g:
            narrowed = [d for d in cands if vendor.norm_gstin(d.fields.vendor_gstin) == g]
            if narrowed:
                return narrowed
        return cands
    return []


def _pick_best(cands: list[Document]) -> Document:
    """Highest confidence wins; deterministic tie-break by doc.id."""
    return max(cands, key=lambda d: (d.confidence, d.id))


# --------------------------------------------------------------------------- #
# 1 · read the return
# --------------------------------------------------------------------------- #
def read_b2b_rows(gstr_path: Path | str,
                  sheet_name: str = DEFAULT_SHEET) -> list[B2BRow]:
    """Parse the B2B sheet once. Read-only — the template is never modified."""
    wb = load_workbook(Path(gstr_path), read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"sheet {sheet_name!r} not found; available: {wb.sheetnames}")
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return []
        headers = {_s(v).lower(): i for i, v in enumerate(header) if v is not None}
        gi = headers.get(GSTIN_HEADER.lower())
        ni = headers.get(INVNO_HEADER.lower())
        if gi is None or ni is None:
            raise ValueError(
                f"could not find '{GSTIN_HEADER}' and '{INVNO_HEADER}' columns in "
                f"sheet {sheet_name!r}; headers seen: {[_s(v) for v in header]}")
        # display-only extras, best effort
        si = next((i for h, i in headers.items() if "trade" in h or "supplier name" in h), None)
        vi = next((i for h, i in headers.items() if "invoice value" in h), None)

        out: list[B2BRow] = []
        for n, values in enumerate(rows_iter, start=2):
            gstin = _s(values[gi]) if gi < len(values) else ""
            invno = _s(values[ni]) if ni < len(values) else ""
            if not gstin and not invno:
                continue  # trailing/blank row
            out.append(B2BRow(
                row=n, gstin=gstin, invoice_no=invno,
                supplier=_s(values[si]) if si is not None and si < len(values) else "",
                value=_s(values[vi]) if vi is not None and vi < len(values) else "",
            ))
        return out
    finally:
        wb.close()


# --------------------------------------------------------------------------- #
# 2 · match (pure)
# --------------------------------------------------------------------------- #
def match(rows: Iterable[B2BRow], invoices: Iterable[Document],
          manual: dict[int, str] | None = None) -> LinkPlan:
    """Resolve every B2B row to an invoice. No I/O — safe to re-run on every edit.

    ``manual`` maps a sheet row number to a doc id a human bound by hand; it wins
    over whatever the matcher would have decided (including over "not found").
    """
    invoices = list(invoices)
    manual = manual or {}
    by_id = {d.id: d for d in invoices}
    by_key, by_invno = _build_index(invoices)

    plan = LinkPlan()
    matched_ids: set[str] = set()

    for r in rows:
        g, n = vendor.norm_gstin(r.gstin), vendor.norm_id(r.invoice_no)

        forced = manual.get(r.row)
        if forced and forced in by_id:
            rm = RowMatch(row=r.row, gstin=r.gstin, invoice_no=r.invoice_no,
                          status=MATCHED, doc_id=forced, manual=True,
                          supplier=r.supplier, value=r.value)
            plan.rows.append(rm)
            plan.by_doc[forced] = rm
            matched_ids.add(forced)
            continue

        cands = _resolve(by_key, by_invno, g, n)
        if not cands:
            plan.rows.append(RowMatch(
                row=r.row, gstin=r.gstin, invoice_no=r.invoice_no,
                status=NOT_FOUND, supplier=r.supplier, value=r.value))
            continue

        # >1 candidate: the same invoice filed twice (a dup — pick the best copy)
        # vs genuinely conflicting GSTINs (a human has to choose).
        gstins = {vendor.norm_gstin(d.fields.vendor_gstin) for d in cands
                  if d.fields.vendor_gstin}
        if len(cands) > 1 and len(gstins) > 1:
            plan.rows.append(RowMatch(
                row=r.row, gstin=r.gstin, invoice_no=r.invoice_no,
                status=AMBIGUOUS, candidates=[d.id for d in cands],
                supplier=r.supplier, value=r.value))
            continue

        doc = _pick_best(cands)
        if len(cands) > 1:
            plan.duplicate_filings += 1
        rm = RowMatch(row=r.row, gstin=r.gstin, invoice_no=r.invoice_no,
                      status=MATCHED, doc_id=doc.id,
                      candidates=[d.id for d in cands] if len(cands) > 1 else [],
                      supplier=r.supplier, value=r.value)
        plan.rows.append(rm)
        plan.by_doc[doc.id] = rm
        matched_ids.add(doc.id)

    plan.unreferenced = [d.id for d in invoices if d.id not in matched_ids]
    return plan


def apply_plan(result: RunResult, plan: LinkPlan) -> None:
    """Stamp the link outcome back onto every invoice Document.

    This is the step the old code was missing: it computed the doc->row mapping
    and threw it away as a count, so nothing downstream (review UI, master xlsx,
    detections.json) could ever see it.
    """
    from ..observability.events import record

    ambiguous_ids: set[str] = set()
    for rm in plan.rows:
        if rm.status == AMBIGUOUS:
            ambiguous_ids.update(rm.candidates)

    for d in result.invoices():
        rm = plan.by_doc.get(d.id)
        if rm is not None:
            d.link_status = LinkStatus.MATCHED
            d.gstr_row = rm.row
            d.gstr_ref = _flat_name(rm.gstin, rm.invoice_no)
            record(d, "link", "matched",
                   f"B2B row {rm.row}" + (" (bound by hand)" if rm.manual else ""),
                   row=rm.row, manual=rm.manual)
        elif d.id in ambiguous_ids:
            d.link_status = LinkStatus.AMBIGUOUS
            d.gstr_row = None
            d.gstr_ref = None
            record(d, "link", "ambiguous",
                   "one of several candidates for a B2B row", severity="warn")
        else:
            d.link_status = LinkStatus.UNREFERENCED
            d.gstr_row = None
            d.gstr_ref = None
            record(d, "link", "unreferenced",
                   "detected invoice satisfies no B2B row", severity="warn")


def suggest(row: B2BRow, invoices: Iterable[Document], limit: int = 5) -> list[tuple[Document, float]]:
    """Near-miss invoices for an unmatched row, best first.

    Reuses the same normalization the matcher keys on (`vendor.norm_id` collapses
    'APEX/22-23/003' and 'APEX-22-23-003'), so a suggestion here is exactly a
    'this nearly matched' — the reviewer can bind it in one click.
    """
    from rapidfuzz import fuzz

    want_n = vendor.norm_id(row.invoice_no)
    want_g = vendor.norm_gstin(row.gstin)
    scored: list[tuple[Document, float]] = []
    for d in invoices:
        cand_ns = [vendor.norm_id(x) for x in
                   (d.fields.invoice_id, d.fields.invoice_no_content) if x]
        if not cand_ns and not d.fields.vendor_gstin:
            continue
        s_num = max((fuzz.ratio(want_n, c) for c in cand_ns if c), default=0.0)
        same_gstin = bool(want_g) and vendor.norm_gstin(d.fields.vendor_gstin) == want_g
        # Weighted so the two signals *share* the scale instead of saturating it:
        # a flat GSTIN bonus on top of a 0-100 number score pinned every same-vendor
        # invoice at 100% and destroyed the ranking. Only an exact number under the
        # right GSTIN reaches 100.
        s = 0.75 * s_num + (25.0 if same_gstin else 0.0)
        if s <= 0:
            continue
        scored.append((d, round(s, 1)))
    scored.sort(key=lambda t: (-t[1], t[0].id))
    return scored[:limit]


# --------------------------------------------------------------------------- #
# 3 · write the linked workbook + flat folder
# --------------------------------------------------------------------------- #
def _flag_cell(cell, text: str) -> None:
    cell.value = text
    cell.fill = _FLAG_FILL


def write_linked(result: RunResult, plan: LinkPlan, gstr_path: Path, store: RunStore,
                 file_source: FileSource | None = None,
                 sheet_name: str = DEFAULT_SHEET,
                 ref_header: str = DEFAULT_REF_HEADER,
                 on_progress: Optional[Callable[[int, int], None]] = None,
                 workers: int = 8) -> LinkReport:
    """Materialize ``plan``: copy the matched PDFs and write ``<stem>_linked.xlsx``.

    ``on_progress(done, total)`` is called as the PDFs land. This is the expensive
    half of linking — on a Drive run every matched row is a network download — so the
    caller runs it on a background thread and shows a real progress bar rather than
    freezing the window for minutes with no sign of life.
    """
    gstr_path = Path(gstr_path)
    file_source = file_source or LocalFileSource()
    by_id = {d.id: d for d in result.invoices()}

    wb = load_workbook(gstr_path)  # read-write, preserves styles/formulas
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"sheet {sheet_name!r} not found; available: {wb.sheetnames}")
    ws = wb[sheet_name]

    headers = {_s(c.value).lower(): c.column for c in ws[1] if c.value is not None}
    ref_col = headers.get(ref_header.lower())
    if not ref_col:  # find-or-append (idempotent on re-run)
        ref_col = ws.max_column + 1
        cell = ws.cell(row=1, column=ref_col, value=ref_header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    # fresh flat folder so re-runs don't leave stale copies
    flat_dir = store.linked_dir
    shutil.rmtree(flat_dir, ignore_errors=True)
    flat_dir.mkdir(parents=True, exist_ok=True)

    counts = plan.counts()
    rep = LinkReport(out_path=store.gstr_linked_path(gstr_path), flat_dir=flat_dir,
                     total=counts["total"], duplicate_filings=plan.duplicate_filings,
                     unreferenced_invoices=counts["unreferenced"])
    used_names: set[str] = set()

    # ---- pass 1: decide every row, serially. No I/O, so the flat-name uniqueness
    # check stays deterministic and openpyxl is only touched from this thread.
    to_copy: list[tuple] = []          # (rm, doc, flat name)
    for rm in plan.rows:
        if rm.status == NOT_FOUND:
            rep.not_found += 1
            rep.unmatched_rows.append((rm.row, rm.gstin, rm.invoice_no))
            _flag_cell(ws.cell(row=rm.row, column=ref_col), "NOT FOUND")
            continue

        if rm.status == AMBIGUOUS:
            rep.ambiguous += 1
            rep.ambiguous_rows.append(
                (rm.row, rm.gstin, rm.invoice_no, len(rm.candidates)))
            _flag_cell(ws.cell(row=rm.row, column=ref_col),
                       f"AMBIGUOUS ({len(rm.candidates)})")
            continue

        doc = by_id.get(rm.doc_id or "")
        if doc is None:  # a bound doc that is no longer an invoice (rejected in review)
            rep.not_found += 1
            rep.unmatched_rows.append((rm.row, rm.gstin, rm.invoice_no))
            _flag_cell(ws.cell(row=rm.row, column=ref_col), "NOT FOUND")
            continue

        name = _flat_name(rm.gstin, rm.invoice_no)
        while name in used_names:  # extremely unlikely; guarantees folder uniqueness
            name = name[:-4] + "-dup.pdf"
        used_names.add(name)
        to_copy.append((rm, doc, name))

    # ---- pass 2: fetch the bytes, in PARALLEL.
    # Bytes come through the FileSource: on a Drive run doc.path is a display string,
    # not a file, so a plain shutil.copy2(doc.path) would fail for every matched row —
    # and each materialize() is a network download. Serially that was the single
    # longest thing the app ever did (hundreds of sequential downloads over a slow
    # link, inside one request, with the UI frozen and no progress). It is network-
    # bound, so threads give real parallelism, and we report progress as they land.
    def _fetch(item) -> tuple:
        rm, doc, name = item
        try:
            local = file_source.materialize(doc)
        except Exception:
            return rm, name, False
        try:
            shutil.copy2(local, flat_dir / name)
            return rm, name, True
        except OSError:
            return rm, name, False
        finally:
            file_source.cleanup(doc, local)

    results: list[tuple] = []
    if to_copy:
        done = 0
        if on_progress:
            on_progress(0, len(to_copy))
        with ThreadPoolExecutor(max_workers=max(1, workers)) as ex:
            for out in ex.map(_fetch, to_copy):
                results.append(out)
                done += 1
                if on_progress:
                    on_progress(done, len(to_copy))

    # ---- pass 3: stamp the outcomes back into the sheet, serially (openpyxl is
    # not thread-safe, so no cell is touched from a worker).
    for rm, name, ok in results:
        cell = ws.cell(row=rm.row, column=ref_col)
        if ok:
            rep.matched += 1
            cell.value = name
        else:
            rep.copy_failed += 1
            _flag_cell(cell, "COPY FAILED")

    _autosize(ws, [ref_header])
    _write_report_sheet(wb, rep, gstr_path, sheet_name)
    wb.save(rep.out_path)
    return rep


def _write_report_sheet(wb, rep: LinkReport, gstr_path: Path, sheet_name: str) -> None:
    if REPORT_SHEET in wb.sheetnames:
        del wb[REPORT_SHEET]
    ws = wb.create_sheet(REPORT_SHEET)
    ws.append(["metric", "value"])
    _style_header(ws, 2)
    for k, v in [
        ("source_template", gstr_path.name),
        ("b2b_sheet", sheet_name),
        ("b2b_rows", rep.total),
        ("matched", rep.matched),
        ("not_found", rep.not_found),
        ("ambiguous", rep.ambiguous),
        ("copy_failed", rep.copy_failed),
        ("duplicate_filings", rep.duplicate_filings),
        ("detected_invoices_unreferenced", rep.unreferenced_invoices),
    ]:
        ws.append([k, v])

    if rep.unmatched_rows:
        ws.append([])
        ws.append(["NOT FOUND — no PDF matched these B2B rows"])
        ws.append(["row", "GSTIN of Supplier", "Invoice Number"])
        for row, gstin, invno in rep.unmatched_rows:
            ws.append([row, gstin, invno])

    if rep.ambiguous_rows:
        ws.append([])
        ws.append(["AMBIGUOUS — multiple conflicting matches"])
        ws.append(["row", "GSTIN of Supplier", "Invoice Number", "candidates"])
        for row, gstin, invno, ncand in rep.ambiguous_rows:
            ws.append([row, gstin, invno, ncand])

    _autosize(ws, ["metric", "value", "col3", "col4"])


# --------------------------------------------------------------------------- #
# convenience: the old one-shot entry point, now composed from the three parts
# --------------------------------------------------------------------------- #
def link_gstr(result: RunResult, gstr_path: Path, store: RunStore,
              sheet_name: str = DEFAULT_SHEET,
              ref_header: str = DEFAULT_REF_HEADER,
              file_source: FileSource | None = None,
              manual: dict[int, str] | None = None) -> LinkReport:
    """Read + match + write in one call (the `link-gstr` CLI and headless paths).

    Manual bindings default to the ones saved in the run dir, so re-linking from
    the CLI doesn't silently discard rows a human already resolved in the review UI.
    """
    rows = read_b2b_rows(gstr_path, sheet_name)
    plan = match(rows, result.invoices(),
                 store.manual_links() if manual is None else manual)
    apply_plan(result, plan)
    store.save_link(plan)
    store.update_meta(gstr_path=str(gstr_path))
    return write_linked(result, plan, gstr_path, store, file_source,
                        sheet_name, ref_header)
